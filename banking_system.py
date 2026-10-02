import openpyxl
import os

file = "banking_system.xlsx"


# Create Excel file
def create_file():
    if not os.path.exists(file):
        wb = openpyxl.Workbook()
        ws = wb.active

        ws.append(["Account Number", "Name", "Mobile Number", "PIN", "Balance"])

        wb.save(file)
        wb.close()


# Create Account
def create_account():
    wb = openpyxl.load_workbook(file)
    ws = wb.active

    print("\n--- Create Account ---")

    name = input("Enter Account Holder Name: ")
    mobile = input("Enter Mobile Number: ")
    pin = input("Enter 4-digit PIN: ")
    deposit = float(input("Enter Initial Deposit: "))

    # Generate account number
    if ws.max_row == 1:
        account_number = 1001
    else:
        account_number = ws.cell(ws.max_row, 1).value + 1

    # Save data in Excel
    ws.append([account_number, name, mobile, pin, deposit])

    wb.save(file)
    wb.close()

    print("\nAccount Created Successfully!")
    print("Your Account Number:", account_number)
    print("Initial Balance:", deposit)


# Find Account
def find_account(account_number):
    wb = openpyxl.load_workbook(file)
    ws = wb.active

    for row in range(2, ws.max_row + 1):

        if ws.cell(row, 1).value == account_number:
            return wb, ws, row

    wb.close()

    return None, None, None


# Deposit Money
def deposit_money():

    print("\n--- Deposit Money ---")

    account_number = int(input("Enter Account Number: "))

    wb, ws, row = find_account(account_number)

    if ws == None:
        print("Account Not Found!")
        return

    pin = input("Enter PIN: ")

    if ws.cell(row, 4).value == pin:

        amount = float(input("Enter Deposit Amount: "))

        if amount > 0:

            balance = ws.cell(row, 5).value

            balance = balance + amount

            ws.cell(row, 5).value = balance

            wb.save(file)
            wb.close()

            print("Deposit Successful!")
            print("Updated Balance:", balance)

        else:
            print("Invalid Amount!")
            wb.close()

    else:
        print("Wrong PIN!")
        wb.close()


# Withdraw Money
def withdraw_money():

    print("\n--- Withdraw Money ---")

    account_number = int(input("Enter Account Number: "))

    wb, ws, row = find_account(account_number)

    if ws == None:
        print("Account Not Found!")
        return

    pin = input("Enter PIN: ")

    if ws.cell(row, 4).value == pin:

        amount = float(input("Enter Withdrawal Amount: "))

        balance = ws.cell(row, 5).value

        if amount <= 0:

            print("Invalid Amount!")

        elif amount > balance:

            print("Insufficient Balance!")

        else:

            balance = balance - amount

            ws.cell(row, 5).value = balance

            wb.save(file)

            print("Withdrawal Successful!")
            print("Updated Balance:", balance)

        wb.close()

    else:

        print("Wrong PIN!")
        wb.close()


# Check Balance
def check_balance():

    print("\n--- Check Balance ---")

    account_number = int(input("Enter Account Number: "))

    wb, ws, row = find_account(account_number)

    if ws == None:

        print("Account Not Found!")
        return

    pin = input("Enter PIN: ")

    if ws.cell(row, 4).value == pin:

        print("\n--- Account Details ---")

        print("Account Number:", ws.cell(row, 1).value)
        print("Account Holder Name:", ws.cell(row, 2).value)
        print("Current Balance:", ws.cell(row, 5).value)

    else:

        print("Wrong PIN!")

    wb.close()


# Main Function
def main():

    create_file()

    while True:

        print("\n==============================")
        print("      SIMPLE BANKING SYSTEM")
        print("==============================")

        print("1. Create Account")
        print("2. Deposit Money")
        print("3. Withdraw Money")
        print("4. Check Balance")
        print("5. Exit")

        choice = input("Enter your choice: ")

        if choice == "1":

            create_account()

        elif choice == "2":

            deposit_money()

        elif choice == "3":

            withdraw_money()

        elif choice == "4":

            check_balance()

        elif choice == "5":

            print("\nThank you for using Simple Banking System!")
            print("Have a nice day!")
            break

        else:

            print("Invalid Choice! Please enter 1 to 5.")


# Start Program
main()